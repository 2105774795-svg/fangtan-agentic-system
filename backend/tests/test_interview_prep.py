"""Synthetic, isolated end-to-end tests; model/ASR boundaries never call the cloud."""
import asyncio
import io
import json
import os
from pathlib import Path
import sys
import tempfile
import unittest
import wave
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
# Import-time migrations must not touch the user's data directory.
_BOOT = tempfile.TemporaryDirectory(prefix='interview-boot-')
os.environ['FANGTAN_DATA_DIR'] = _BOOT.name
os.environ['FANGTAN_AUTH_ENABLED'] = 'false'
import app
import store
import config
import worker
import job_queue
from interview_prep import domain as d, ai
from fastapi.testclient import TestClient
ORIGINAL_GENERATE = ai.generate


def model_fixture(kind, data):
    if kind=='interview_question_analysis':
        return {'intent':'考查岗位动机与经历匹配','answer_outline':['说明岗位理解','连接真实经历'],
                'previous_answer_feedback':['上次回答需要更具体'], 'followup_angles':['追问个人行动'],
                'preparation_actions':['准备项目证据'],'missing_information':[]}
    if kind == 'interview_prepare':
        return {'summary':'根据真实项目经历准备面试', 'preparation_tips':['用事实说明个人行动'], 'questions':[{'question':'请介绍你负责的项目和个人贡献？','target':'个人贡献'}, {'question':'你作出了哪些重要取舍？','target':'决策依据'}]}
    answer = [a for t in data['turns'] for a in t['answers']][-1]
    finding = {'dimension':'个人贡献', 'observation':'回答提供了个人行动', 'suggestion':'补充结果依据', 'answer_id':answer['id'], 'quote':answer['transcript'][:10], 'segment_ids':[]}
    if kind == 'interview_next_turn':
        return {'action':'follow_up', 'question':{'question':'你刚才提到的结果是如何验证的？','target':'结果依据'},'feedback':[finding]}
    return {'summary':'本轮已能说明行动，可进一步补充证据。', 'findings':[finding], 'retry_questions':[{'question':'请用证据说明你的项目结果？','target':'结果依据'}]}


class InterviewPreparationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='interview-test-')
        self.root = Path(self.temp.name)
        self.who = 11
        async def user(request): return {'id':self.who,'name':'合成用户'}
        self.patches = [patch.object(store,'DB_PATH',self.root/'test.db'), patch.object(store,'DATA_DIR',self.root),
                        patch.object(config,'DATA_DIR',self.root),patch.object(config,'UPLOAD_DIR',self.root/'uploads'),patch.object(config,'REPORT_DIR',self.root/'reports'),
                        patch.object(config,'AUTH_ENABLED',True),patch.object(config,'KIMI_API_KEY','synthetic-only'),
                        patch('app.employee_auth.current_user',side_effect=user), patch.object(ai,'generate',side_effect=model_fixture)]
        for p in self.patches:p.start()
        store.init();job_queue.init();d.init()
        config.UPLOAD_DIR.mkdir();config.REPORT_DIR.mkdir()
        self.client=TestClient(app.app)
        self.cfg={'template':'job','goal':'应聘产品经理，练习个人贡献','practice_mode':'coaching','question_count':2,'max_followups':1,'followup_mode':'fixed','include_candidate_questions':False,'duration_minutes':20,'language':'中文'}
        self.project=self.client.post('/api/interview-prep/projects',json={'title':'合成面试','request_id':'create-123456','config':self.cfg}).json()

    def tearDown(self):
        self.client.close()
        for p in reversed(self.patches):p.stop()
        self.temp.cleanup()

    def refresh(self):
        response=self.client.get('/api/interview-prep/projects/'+self.project['id'])
        self.assertEqual(response.status_code,200,response.text)
        self.project=response.json();return self.project

    def command(self, action, expected=200, **fields):
        body={'revision':self.project['revision'],'request_id':d.uuid.uuid4().hex,'action':action,**fields}
        response=self.client.post('/api/interview-prep/projects/'+self.project['id']+'/commands',json=body)
        self.assertEqual(response.status_code,expected,response.text)
        if response.status_code==200:self.project=response.json()
        return response

    def run_task(self):
        self.assertTrue(worker.run_once(handlers=ai.HANDLERS,worker_id='synthetic-test'))
        self.refresh()
        job=self.project['practices'][0]['job'] if self.project['practices'] else self.project['job']
        self.assertEqual(job['status'],'ready',job)

    def prepare(self):
        self.command('material',kind='resume',title='合成简历',content='我负责项目需求分析，设计流程并组织验证。')
        self.command('prepare');self.run_task()
        self.command('start')
        return self.project['practices'][0]

    def answer(self, practice=None):
        practice=practice or self.project['practices'][0]
        turn=next(t for t in practice['turns'] if t['status']=='waiting')
        self.command('answer',practice_id=practice['id'],turn_id=turn['id'])
        turn=next(t for t in self.project['practices'][0]['turns'] if t['id']==turn['id'])
        return turn['answers'][-1]

    def upload(self, answer, text='我设计了流程并通过测试验证结果。', chunk='chunk-synthetic-001', fail=False):
        audio=io.BytesIO()
        with wave.open(audio,'wb') as f:
            f.setnchannels(1);f.setsampwidth(2);f.setframerate(16000);f.writeframes(bytes(32000))
        with patch('app.asr_dashscope.transcribe',side_effect=RuntimeError('synthetic outage') if fail else None,return_value=text), patch('app.coverage_ai.schedule') as coverage:
            response=self.client.post('/api/interview-prep/answers/'+answer['id']+'/audio',data={'chunk_id':chunk},files={'audio':('answer.wav',audio.getvalue(),'audio/wav')})
            self.assertEqual(response.status_code,502 if fail else 200,response.text)
            coverage.assert_not_called()
        return chunk

    def complete(self,answer,chunks,expected=200):
        return self.command('complete_answer',expected=expected,practice_id=self.project['practices'][0]['id'],answer_id=answer['id'],chunk_ids=chunks)

    def test_full_recording_followup_review_correction_and_retry(self):
        self.prepare();answer=self.answer();chunk=self.upload(answer);self.complete(answer,[chunk]);self.run_task()
        p=self.project['practices'][0]
        self.assertIsNotNone(p['turns'][1]['parent_id'])
        self.assertEqual(self.client.get('/api/interview-prep/answers/'+answer['id']+'/recording').status_code,200)
        second=self.answer();chunk=self.upload(second,chunk='chunk-synthetic-002');self.complete(second,[chunk]);self.run_task()
        # One follow-up per main question: the next turn must use the plan.
        self.assertIsNone(self.project['practices'][0]['turns'][-1]['parent_id'])
        self.command('finish',practice_id=p['id']);self.command('review',practice_id=p['id']);self.run_task()
        self.assertTrue(self.project['practices'][0]['review']['findings'])
        self.command('correct',practice_id=p['id'],answer_id=answer['id'],content='人工校对：我负责需求分析和测试。')
        self.assertTrue(self.project['practices'][0]['review_stale'])
        self.command('retry_practice',practice_id=p['id'],expected=409)
        self.command('review',practice_id=p['id']);self.run_task()
        self.command('retry_practice',practice_id=p['id'])
        self.assertEqual(self.project['practices'][0]['source_practice_id'],p['id'])
        self.assertEqual(self.project['practices'][0]['turns'][0]['target'],'结果依据')

    def test_manifest_and_no_speech_cannot_advance(self):
        self.prepare();answer=self.answer();self.complete(answer,[],expected=409)
        c=self.upload(answer,text='');self.complete(answer,[c],expected=422)
        self.assertEqual(self.refresh()['practices'][0]['turns'][0]['status'],'waiting')
        self.upload(answer,text='补录后有回答。',chunk='chunk-synthetic-002')
        self.complete(answer,['chunk-synthetic-002'],expected=409)
        self.complete(answer,[c,'chunk-synthetic-002'])

    def test_asr_retry_and_duplicate_chunk_preserve_exactly_one_utterance(self):
        self.prepare();answer=self.answer();c=self.upload(answer,fail=True)
        self.complete(answer,[c],expected=409)
        self.upload(answer);self.upload(answer)
        self.complete(answer,[c])
        with store.get() as db:
            self.assertEqual(db.execute('SELECT COUNT(*) FROM utterances WHERE session_id=?',(answer['session_id'],)).fetchone()[0],1)
            self.assertEqual(db.execute('SELECT COUNT(*) FROM reports WHERE session_id=?',(answer['session_id'],)).fetchone()[0],0)
        self.assertEqual(self.client.get('/api/sessions/'+answer['session_id']).status_code,404)
        self.assertEqual(store.list_sessions()['total'],0)

    def test_idempotent_completion_and_stale_revision(self):
        self.prepare();a=self.answer();c=self.upload(a)
        body={'action':'complete_answer','revision':self.project['revision'],'request_id':'same-complete-123','practice_id':self.project['practices'][0]['id'],'answer_id':a['id'],'chunk_ids':[c]}
        url='/api/interview-prep/projects/'+self.project['id']+'/commands'
        for _ in range(2): self.assertEqual(self.client.post(url,json=body).status_code,200)
        with store.get() as db:self.assertEqual(db.execute("SELECT COUNT(*) FROM tasks WHERE kind='interview_next_turn'").fetchone()[0],1)
        body.update(action='next',request_id='new-request-123')
        self.assertEqual(self.client.post(url,json=body).status_code,409)

    def test_owner_isolation_and_deletion(self):
        self.prepare();a=self.answer();self.upload(a)
        pid=self.project['id'];self.who=22
        self.assertEqual(self.client.get('/api/interview-prep/projects').json(),[])
        self.assertEqual(self.client.get('/api/interview-prep/projects/'+pid).status_code,404)
        self.assertEqual(self.client.get('/api/interview-prep/answers/'+a['id']+'/recording').status_code,404)
        self.who=11
        self.assertEqual(self.client.delete('/api/interview-prep/projects/'+pid).status_code,200)
        self.assertEqual(self.client.get('/api/interview-prep/projects/'+pid).status_code,404)
        self.assertEqual(self.client.get('/api/interview-prep/answers/'+a['id']+'/recording').status_code,404)

    def test_simulation_hides_feedback_until_finished(self):
        self.command('configure',config={**self.cfg,'practice_mode':'simulation'})
        self.prepare();a=self.answer();self.complete(a,[self.upload(a)]);self.run_task()
        self.assertIsNone(self.project['practices'][0]['turns'][0]['feedback'])
        self.command('finish',practice_id=self.project['practices'][0]['id'])
        self.assertTrue(self.project['practices'][0]['turns'][0]['feedback'])

    def test_material_change_does_not_change_frozen_practice(self):
        p=self.prepare()
        self.command('material',title='新岗位',kind='job_description',content='新岗位要求新的能力。')
        self.assertTrue(self.project['plan_stale'])
        self.assertEqual(len(self.project['practices'][0]['material_versions']),1)
        self.assertEqual(len(self.project['materials']),2)

    def test_cancellation_discards_late_model_result(self):
        self.command('material',title='简历',content='合成的经历。');self.command('prepare')
        def cancel_during_generation(kind,data):
            self.refresh()
            self.command('material',title='新材料',content='另一个新事实。')
            return model_fixture(kind,data)
        with patch.object(ai,'generate',side_effect=cancel_during_generation):worker.run_once(handlers=ai.HANDLERS)
        self.refresh();self.assertIsNone(self.project['plan'])
        self.assertEqual(self.project['job']['status'],'cancelled')

    def test_evidence_validation_rejects_fabricated_quote_or_segment(self):
        data={'turns':[{'answers':[{'id':'a','transcript':'我做了测试','segments':[{'id':1,'text':'我做了测试'}]}]}]}
        f={'answer_id':'a','quote':'不存在的成绩','segment_ids':[]}
        with self.assertRaises(ValueError):ai.validate_evidence([f],data)
        f.update(quote='我做了测试',segment_ids=[999])
        with self.assertRaises(ValueError):ai.validate_evidence([f],data)
        f.update(segment_ids=[1]);ai.validate_evidence([f],data)

    def test_upload_preparation_material_is_not_an_interview(self):
        response=self.client.post('/api/interview-prep/projects/'+self.project['id']+'/materials',data={'revision':self.project['revision'],'request_id':'upload-123456','kind':'resume'},files={'file':('resume.txt','合成简历，负责流程设计。'.encode(),'text/plain')})
        self.assertEqual(response.status_code,200,response.text)
        self.assertEqual(store.list_sessions()['total'],0)
        bad=self.client.post('/api/interview-prep/projects/'+self.project['id']+'/materials',data={'revision':response.json()['revision'],'request_id':'upload-987654','kind':'paper'},files={'file':('paper.pdf',b'not a pdf','application/pdf')})
        self.assertEqual(bad.status_code,422)


    def test_material_revision_preserves_history_and_frozen_input(self):
        self.prepare()
        original=self.project['materials'][0]
        self.command('revise_material',material_id=original['id'],content='修订：我负责项目设计和验证。')
        updated=self.project['materials'][0]
        self.assertEqual(updated['version'],2)
        self.assertEqual(json.loads(updated['history'])[0]['content'],original['content'])
        self.assertEqual(self.project['practices'][0]['material_versions'][0]['version'],1)
        self.assertTrue(self.project['plan_stale'])

    def test_pptx_and_text_pdf_parsers(self):
        from interview_prep.api import extract_material
        from pptx import Presentation
        presentation=Presentation();slide=presentation.slides.add_slide(presentation.slide_layouts[1]);slide.shapes.title.text='Synthetic thesis';slide.placeholders[1].text='Research method and evidence'
        buffer=io.BytesIO();presentation.save(buffer)
        self.assertIn('Research method',extract_material(buffer.getvalue(),'slides.pptx'))
        # A minimal, real text PDF with one page, no downloaded materials.
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        writer=PdfWriter();page=writer.add_blank_page(width=300,height=300)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        stream=DecodedStreamObject();stream.set_data(b'BT /F1 12 Tf 20 250 Td (Synthetic research method) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(stream)
        output=io.BytesIO();writer.write(output)
        self.assertIn('Synthetic research method',extract_material(output.getvalue(),'paper.pdf'))
        blank=PdfWriter();blank.add_blank_page(width=100,height=100);output=io.BytesIO();blank.write(output)
        from fastapi import HTTPException
        with self.assertRaises(HTTPException) as rejected:extract_material(output.getvalue(),'scan.pdf')
        self.assertIn('OCR',rejected.exception.detail)

    def test_model_schema_and_evidence_validation_at_cloud_boundary(self):
        data={'config':self.cfg,'materials':[]}
        with patch.object(ai,'generate',ORIGINAL_GENERATE), patch('interview_prep.ai.kimi._chat',return_value=json.dumps(model_fixture('interview_prepare',data))):
            result=ai.generate('interview_prepare',data)
            self.assertEqual(len(result['questions']),2)
        data['turns']=[{'answers':[{'id':'answer','transcript':'真实回答','segments':[]}]}]
        invalid=model_fixture('interview_review',data);invalid['findings'][0]['quote']='虚构原话'
        with patch.object(ai,'generate',ORIGINAL_GENERATE), patch('interview_prep.ai.kimi._chat',return_value=json.dumps(invalid)):
            with self.assertRaises(ValueError):ai.generate('interview_review',data)

    def test_private_responses_and_validation_do_not_echo_material(self):
        response=self.client.get('/api/interview-prep/projects/'+self.project['id'])
        self.assertEqual(response.headers['cache-control'],'private, no-store')
        bad=self.client.post('/api/interview-prep/projects',json={'title':'private phrase','config':{'goal':'private phrase'}})
        self.assertEqual(bad.status_code,422)
        self.assertNotIn('private phrase',bad.text)


    def test_completed_answer_survives_missing_model_configuration(self):
        self.prepare();answer=self.answer();chunk=self.upload(answer)
        with patch.object(config,'KIMI_API_KEY',''):
            self.complete(answer,[chunk])
        self.assertEqual(self.project['practices'][0]['turns'][0]['answers'][0]['status'],'ready')
        self.assertEqual(self.project['practices'][0]['job']['status'],'queued')

    def test_authenticated_gate_rejects_anonymous_access(self):
        from fastapi import HTTPException
        async def reject(request):raise HTTPException(401,'请先登录')
        with patch('app.employee_auth.current_user',side_effect=reject):
            self.assertEqual(self.client.get('/api/interview-prep/projects').status_code,401)
            self.assertEqual(self.client.get('/api/interview-prep/projects/'+self.project['id']).status_code,401)


    def test_delete_material_invalidates_plan_but_preserves_practice_snapshot(self):
        self.prepare();material=self.project['materials'][0]
        self.command('delete_material',material_id=material['id'])
        self.assertEqual(self.project['materials'],[])
        self.assertTrue(self.project['plan_stale'])
        self.assertEqual(self.project['practices'][0]['material_versions'][0]['id'],material['id'])
        self.command('delete_material',material_id=material['id'],expected=404)
        # A goal-only plan is allowed; deleted personal materials are not reused.
        self.command('prepare')
        with store.get() as db:
            context = d.task_input(db, d.owned(db,self.project['id'],self.who))
        self.assertEqual(context['materials'],[])

    def test_dialogue_progress_persists_and_goal_only_plan_works(self):
        self.command('configure',config={**self.cfg,'preparation_step':'closing'})
        self.assertEqual(self.refresh()['config']['preparation_step'],'closing')
        self.command('configure',config={**self.cfg,'preparation_step':'ready','include_candidate_questions':True})
        self.command('prepare');self.run_task()
        self.assertEqual(self.project['materials'],[])
        self.assertTrue(self.project['plan']['questions'])
        self.assertTrue(self.refresh()['config']['include_candidate_questions'])

    def test_review_addition_enters_next_practice_without_regenerating_plan(self):
        self.prepare();answer=self.answer();chunk=self.upload(answer);self.complete(answer,[chunk]);self.run_task()
        pid=self.project['practices'][0]['id']
        self.command('finish',practice_id=pid);self.command('review',practice_id=pid);self.run_task()
        self.command('add_question',question='你为什么从上家公司离职？')
        self.assertTrue(self.project['plan_stale'])
        self.command('retry_practice',practice_id=pid)
        practice=self.project['practices'][0]
        self.assertEqual(practice['turns'][0]['question'],'你为什么从上家公司离职？')
        with store.get() as db:
            frozen=json.loads(db.execute('SELECT snapshot FROM interview_practices WHERE id=?',(practice['id'],)).fetchone()['snapshot'])
        self.assertIn('请用证据说明你的项目结果？',[q['question'] for q in frozen['plan']['questions']])
        self.assertEqual(practice['source_practice_id'],pid)
        self.assertTrue(self.project['practices'][1]['review'])

    def test_delete_material_cancels_inflight_preparation_and_checks_owner(self):
        self.command('material',title='误传文件',content='合成误传材料')
        material=self.project['materials'][0]
        self.command('prepare')
        self.who=22
        self.command('delete_material',material_id=material['id'],expected=404)
        self.who=11
        self.command('delete_material',material_id=material['id'])
        self.assertEqual(self.project['job']['status'],'cancelled')
        self.assertEqual(self.project['materials'],[])

    def test_answer_preview_requires_confirmation_before_model_task(self):
        self.prepare();answer=self.answer();chunk=self.upload(answer)
        pid=self.project['practices'][0]['id']
        self.command('prepare_answer',practice_id=pid,answer_id=answer['id'],chunk_ids=[chunk])
        preview=self.project['practices'][0]['turns'][0]['answers'][0]
        self.assertEqual(preview['status'],'confirming')
        self.assertEqual(preview['transcript'],preview['original_transcript'])
        self.assertEqual(self.client.get('/api/interview-prep/answers/'+answer['id']+'/recording').status_code,200)
        with store.get() as db:
            self.assertEqual(db.execute("SELECT COUNT(*) FROM tasks WHERE kind='interview_next_turn'").fetchone()[0],0)
        self.command('finish',practice_id=pid,expected=409)
        self.command('answer',practice_id=pid,turn_id=self.project['practices'][0]['turns'][0]['id'],expected=409)
        self.command('confirm_answer',practice_id=pid,answer_id=answer['id'],content='我负责设计流程并测试验证结果。')
        saved=self.project['practices'][0]['turns'][0]['answers'][0]
        self.assertEqual(saved['status'],'ready');self.assertEqual(saved['version'],2)
        self.assertNotEqual(saved['original_transcript'],saved['transcript'])
        self.run_task()

    def test_preview_can_be_discarded_and_re_recorded(self):
        p=self.prepare();answer=self.answer();chunk=self.upload(answer)
        self.command('prepare_answer',practice_id=p['id'],answer_id=answer['id'],chunk_ids=[chunk])
        self.command('discard_answer',practice_id=p['id'],answer_id=answer['id'])
        another=self.answer();self.assertNotEqual(another['id'],answer['id'])

    def test_typography_alignment_preserves_exact_original_quote(self):
        data={'turns':[{'answers':[{'id':'a','transcript':'我负责需求分析，随后组织了测试。','segments':[]}]}]}
        f={'answer_id':'a','quote':'我负责需求分析, 随后组织了测试。','segment_ids':[]}
        ai.validate_evidence([f],data)
        self.assertEqual(f['quote'],'我负责需求分析，随后组织了测试。')
        f['quote']='我负责需求分析，并且提高了业绩。'
        with self.assertRaises(ValueError):ai.validate_evidence([f],data)

    def test_evidence_catalog_resolves_to_source_not_model_paraphrase(self):
        data={'turns':[{'answers':[{'id':'a','transcript':'我负责调研并设计流程。','segments':[{'id':1,'text':'我负责调研并设计流程。'}]}]}]}
        catalog=ai.evidence_catalog(data,current_only=True)
        f={'answer_id':'a','evidence_id':'E1','quote':'模型改写的句子','segment_ids':[999]}
        ai.resolve_evidence([f],catalog);ai.validate_evidence([f],data)
        self.assertEqual(f['quote'],'我负责调研并设计流程。');self.assertEqual(f['segment_ids'],[1])

    def test_failed_quote_repair_does_not_block_next_question(self):
        self.prepare();a=self.answer();self.complete(a,[self.upload(a)])
        invalid={'action':'follow_up','question':{'question':'这个不存在的结果是什么？','target':'结果'},'feedback':[{'dimension':'结果','observation':'无依据','suggestion':'提供证据','answer_id':a['id'],'quote':'从未说过的虚构数字','segment_ids':[]}]}
        with patch.object(ai,'generate',ORIGINAL_GENERATE),patch('interview_prep.ai.kimi._chat',return_value=json.dumps(invalid)) as call:
            self.run_task()
            self.assertEqual(call.call_count,2)
        p=self.project['practices'][0]
        self.assertEqual(p['turns'][-1]['question'],'你作出了哪些重要取舍？')
        self.assertEqual(p['turns'][0]['feedback'],[])
        self.assertIn('未通过原话核对',p['turns'][0]['feedback_notice'])

    def test_quote_repair_uses_valid_second_response(self):
        data={'config':self.cfg,'turns':[{'answers':[{'id':'a','transcript':'我完成了流程设计并组织测试。','segments':[]}]}]}
        invalid={'action':'next','feedback':[{'dimension':'行动','observation':'说明行动','suggestion':'补充结果','answer_id':'a','quote':'虚构','segment_ids':[]}]}
        corrected=json.loads(json.dumps(invalid));corrected['feedback'][0]['quote']='我完成了流程设计并组织测试。'
        with patch.object(ai,'generate',ORIGINAL_GENERATE),patch('interview_prep.ai.kimi._chat',side_effect=[json.dumps(invalid),json.dumps(corrected)]) as call:
            result=ai.generate('interview_next_turn',data)
            self.assertEqual(call.call_count,2)
            self.assertEqual(result['feedback'][0]['quote'],corrected['feedback'][0]['quote'])


    def test_historical_question_analysis_and_exact_repractice(self):
        self.command('add_question',question='你为什么选择我们公司？',previous_answer='因为贵公司很优秀。',question_context='产品岗位一面')
        q=self.project['question_bank'][0]
        self.command('analyze_question',question_id=q['id']);self.run_task()
        self.assertIn('匹配',self.project['question_bank'][0]['analysis']['intent'])
        self.command('practice_question',question_id=q['id'])
        self.assertEqual(self.project['practices'][0]['turns'][0]['question'],q['question'])
        self.assertEqual(self.project['practices'][0]['turns'][0]['answers'],[])
        with store.get() as db:
            frozen=json.loads(db.execute('SELECT snapshot FROM interview_practices WHERE id=?',(self.project['practices'][0]['id'],)).fetchone()[0])
            self.assertEqual(frozen['question_bank'][0]['previous_answer'],'因为贵公司很优秀。')

    def test_historical_questions_take_priority_in_plan(self):
        self.command('add_question',question='上次问到的真实问题是什么？')
        self.command('prepare');self.run_task()
        self.assertEqual(self.project['plan']['questions'][0]['question'],'上次问到的真实问题是什么？')

    def test_company_and_role_context_is_frozen_and_analysis_becomes_stale(self):
        self.command('configure',config={**self.cfg,'company_name':'合成公司','role_name':'产品经理',
                    'company_background':'提供企业预约产品','job_requirements':'负责需求分析','interview_stage':'业务初面'})
        self.command('add_question',question='你了解我们的业务吗？')
        q=self.project['question_bank'][0]
        self.command('analyze_question',question_id=q['id']);self.run_task()
        self.command('prepare');self.run_task();self.command('start')
        self.assertEqual(self.project['practices'][0]['config']['company_name'],'合成公司')
        self.command('configure',config={**self.cfg,'company_name':'另一家合成公司'})
        self.assertTrue(self.project['question_bank'][0]['analysis_stale'])
        self.assertEqual(self.project['practices'][0]['config']['company_name'],'合成公司')

    def test_adaptive_followup_can_exceed_former_numeric_limit(self):
        self.command('configure',config={**self.cfg,'followup_mode':'adaptive','max_followups':0})
        self.prepare();a=self.answer();self.complete(a,[self.upload(a)]);self.run_task()
        self.assertIsNotNone(self.project['practices'][0]['turns'][-1]['parent_id'])
        def distinct_followup(kind,data):
            result=model_fixture(kind,data)
            result['question']={'question':'你选择这个验证方法的原因是什么？','target':'决策依据'}
            result['follow_up_reason']='刚才回答尚未解释验证方法的选择'
            return result
        a=self.answer();self.complete(a,[self.upload(a,chunk='chunk-adaptive-002')])
        with patch.object(ai,'generate',side_effect=distinct_followup):self.run_task()
        self.assertIsNotNone(self.project['practices'][0]['turns'][-1]['parent_id'])
        self.assertIn('验证方法',self.project['practices'][0]['turns'][-1]['decision_reason'])

    def test_repeated_adaptive_followup_moves_to_next_plan_question(self):
        self.command('configure',config={**self.cfg,'followup_mode':'adaptive'})
        self.prepare();a=self.answer();self.complete(a,[self.upload(a)]);self.run_task()
        a=self.answer();self.complete(a,[self.upload(a,chunk='chunk-repeat-002')]);self.run_task()
        self.assertEqual(self.project['practices'][0]['turns'][-1]['question'],'你作出了哪些重要取舍？')
        self.assertIsNone(self.project['practices'][0]['turns'][-1]['parent_id'])

    def test_candidate_questions_manual_entry_evaluation_and_end(self):
        self.command('configure',config={**self.cfg,'include_candidate_questions':True})
        p=self.prepare();self.command('candidate_questions',practice_id=p['id'])
        closing=self.project['practices'][0]['turns'][-1]
        self.assertEqual(closing['turn_kind'],'candidate_questions')
        a=self.answer();self.complete(a,[self.upload(a,text='这个岗位入职前三个月最重要的目标是什么？')])
        review={'summary':'问题聚焦岗位预期，有助于双向选择。','strengths':['关注岗位目标'],
                'improvements':['可以进一步问评价方式'],
                'recommended_questions':[{'question':'团队会如何评价这个岗位前三个月的成果？','why':'了解成功标准','timing':'业务主管面试收尾'}]}
        def closing_fixture(kind,data):
            result=model_fixture(kind,data);result.update(action='finish',candidate_questions_review=review);return result
        with patch.object(ai,'generate',side_effect=closing_fixture):self.run_task()
        result=self.project['practices'][0]
        self.assertEqual(result['status'],'done')
        self.assertIn('双向选择',result['candidate_review']['summary'])
        self.assertEqual(result['turns'][-1]['answers'][0]['status'],'ready')

    def test_normal_completion_enters_candidate_questions_automatically(self):
        self.command('configure',config={**self.cfg,'max_followups':0,'include_candidate_questions':True})
        self.prepare()
        for i in range(2):
            a=self.answer();self.complete(a,[self.upload(a,chunk=f'chunk-complete-{i:03}')]);self.run_task()
        self.assertEqual(self.project['practices'][0]['status'],'active')
        self.assertEqual(self.project['practices'][0]['turns'][-1]['turn_kind'],'candidate_questions')

    def test_question_bank_owner_check_and_deletion(self):
        self.command('add_question',question='如何处理用户需求冲突？')
        q=self.project['question_bank'][0]
        self.who=22;self.command('delete_question',question_id=q['id'],expected=404)
        self.who=11;self.command('delete_question',question_id=q['id'])
        self.assertEqual(self.project['question_bank'],[])

    def test_thesis_does_not_get_job_candidate_questions(self):
        self.command('configure',config={**self.cfg,'template':'thesis','include_candidate_questions':True,'max_followups':0})
        self.prepare()
        for i in range(2):
            a=self.answer();self.complete(a,[self.upload(a,chunk=f'chunk-thesis-{i:03}')]);self.run_task()
        self.assertEqual(self.project['practices'][0]['status'],'done')
        self.assertFalse(any(t['turn_kind']=='candidate_questions' for t in self.project['practices'][0]['turns']))

if __name__=='__main__':unittest.main()
