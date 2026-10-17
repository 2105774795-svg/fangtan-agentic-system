"""Interview APIs: every request, stream, audio and task checks project ownership."""
import asyncio
import io
import json
import tempfile
import threading
import zipfile
from pathlib import Path
from typing import Literal
from fastapi import APIRouter, Request, HTTPException, UploadFile, File, Form
from fastapi.responses import StreamingResponse, FileResponse
from starlette.concurrency import run_in_threadpool
from pydantic import BaseModel, Field, ConfigDict
import store
import importer
import config
from . import domain as d

router = APIRouter(prefix='/api/interview-prep', tags=['interview-preparation'])

class PreparationConfig(BaseModel):
    model_config = ConfigDict(extra='forbid')
    template: Literal['job', 'thesis'] = 'job'
    goal: str = Field(min_length=1, max_length=2000)
    preparation_step: Literal['goal', 'background', 'closing', 'ready'] = 'goal'
    language: Literal['中文', 'English'] = '中文'
    practice_mode: Literal['coaching', 'simulation'] = 'coaching'
    question_count: int = Field(default=5, ge=2, le=12)
    followup_mode: Literal['adaptive','fixed'] = 'adaptive'
    include_candidate_questions: bool = True
    company_name: str = Field(default='', max_length=200)
    role_name: str = Field(default='', max_length=200)
    company_background: str = Field(default='', max_length=6000)
    job_requirements: str = Field(default='', max_length=6000)
    background_source: str = Field(default='', max_length=1000)
    interview_stage: str = Field(default='', max_length=200)
    max_followups: int = Field(default=1, ge=0, le=3)
    duration_minutes: int = Field(default=20, ge=3, le=90)

class Create(BaseModel):
    title: str = Field(min_length=1, max_length=120)
    request_id: str = Field(min_length=8, max_length=80)
    config: PreparationConfig

class Command(BaseModel):
    model_config = ConfigDict(extra='forbid')
    action: Literal['add_question','delete_question','analyze_question','practice_question','candidate_questions','configure','material','revise_material','delete_material','prepare','start','answer','prepare_answer','confirm_answer','complete_answer','discard_answer','finish','next','review','retry_practice','correct','cancel']
    revision: int = Field(ge=1)
    request_id: str = Field(min_length=8, max_length=80)
    config: PreparationConfig | None = None
    title: str = Field(default='', max_length=120)
    kind: Literal['resume','company_background','job_description','paper','presentation','background'] = 'background'
    content: str = Field(default='', max_length=100000)
    practice_id: str | None = None
    turn_id: str | None = None
    answer_id: str | None = None
    material_id: str | None = None
    question_id: str | None = None
    question: str = Field(default='', max_length=500)
    previous_answer: str = Field(default='', max_length=10000)
    question_context: str = Field(default='', max_length=2000)
    chunk_ids: list[str] = Field(default_factory=list, max_length=1000)


def owner(request):
    user = getattr(request.state, 'user', None)
    if not user: raise HTTPException(401, '请先登录，或在本机显式启用单人模式')
    return int(user['id'])

@router.get('/capabilities')
def capabilities(request: Request):
    owner(request)
    return {'templates': d.TEMPLATES, 'model_configured': bool(config.KIMI_API_KEY), 'transcription_configured': bool(config.DASHSCOPE_API_KEY), 'local_mode': not config.AUTH_ENABLED}

@router.get('/projects')
def projects(request: Request):
    with store.get() as db:
        return [dict(r) for r in db.execute('''SELECT id,title,revision,updated_at,
                CASE WHEN EXISTS(SELECT 1 FROM interview_practices p WHERE p.project_id=interview_preparations.id AND p.status='active') THEN 3
                     WHEN EXISTS(SELECT 1 FROM interview_practices p WHERE p.project_id=interview_preparations.id AND p.review IS NOT NULL) THEN 5
                     WHEN plan IS NOT NULL THEN 2 ELSE 0 END AS stage
                FROM interview_preparations WHERE owner_id=? AND deleted_at=0 ORDER BY updated_at DESC''', (owner(request),))]

@router.post('/projects', status_code=201)
def create(body: Create, request: Request):
    if not body.title.strip() or not body.config.goal.strip(): raise HTTPException(422, '请填写面试目标和项目名称')
    return d.create(owner(request), body.title.strip(), body.request_id, body.config.model_dump())

@router.get('/projects/{pid}')
def project(pid: str, request: Request):
    with store.get() as db: return d.snapshot(db, pid, owner(request))

@router.post('/projects/{pid}/commands')
def command(pid: str, body: Command, request: Request):
    if body.action == 'configure' and body.config is None: raise HTTPException(422, '缺少准备配置')
    # Serialize finalization with the common session upload lock, then verify receipts
    # under the write transaction. A late chunk can never enter a completed answer.
    if body.action in {'complete_answer', 'prepare_answer', 'discard_answer'}:
        from app import CHUNK_LOCKS
        with store.get() as db:
            d.owned(db, pid, owner(request))
            row = db.execute('''SELECT a.session_id FROM interview_answers a JOIN interview_turns t ON a.turn_id=t.id
                              JOIN interview_practices p ON t.practice_id=p.id WHERE a.id=? AND p.project_id=?''', (body.answer_id, pid)).fetchone()
            if not row: raise HTTPException(404, '回答不存在')
        with CHUNK_LOCKS.setdefault(row['session_id'], threading.Lock()):
            result = d.command(pid, owner(request), body.model_dump())
            from app import _assemble_live_recording
            _assemble_live_recording(row['session_id'])
            return result
    return d.command(pid, owner(request), body.model_dump())

@router.delete('/projects/{pid}')
def delete(pid: str, request: Request):
    with store.get() as db:
        db.execute('BEGIN IMMEDIATE')
        p = d.owned(db, pid, owner(request))
        db.execute('UPDATE interview_preparations SET deleted_at=? WHERE id=?', (d.time.time(), pid))
        if p.get('workspace_project_id'):
            db.execute('UPDATE research_workspace_runs SET deleted_at=? WHERE project_id=?', (d.time.time(), p['workspace_project_id']))
            db.execute("UPDATE tasks SET status='cancelled',updated_at=?,finished_at=? WHERE resource_id=? AND status IN ('queued','running','retry')", (d.time.time(), d.time.time(), p['workspace_project_id']))
            db.execute('DELETE FROM research_workspace_layouts WHERE project_id=?', (p['workspace_project_id'],))
            for table in ('workspace_chat_messages', 'workspace_chat_runs', 'workspace_events'):
                db.execute(f'DELETE FROM {table} WHERE project_id=?', (p['workspace_project_id'],))
        db.execute("UPDATE tasks SET status='cancelled',updated_at=?,finished_at=? WHERE resource_id=? AND kind LIKE 'interview_%' AND status IN ('queued','running','retry')", (d.time.time(), d.time.time(), pid))
    return {'deleted': True, 'retention': '原始录音保留在本机数据目录，删除项目后不再通过本模块访问'}


def answer_access(db, aid, who):
    row = db.execute('''SELECT a.*,p.project_id FROM interview_answers a JOIN interview_turns t ON a.turn_id=t.id
                       JOIN interview_practices p ON t.practice_id=p.id WHERE a.id=?''', (aid,)).fetchone()
    if not row: raise HTTPException(404, '回答不存在')
    d.owned(db, row['project_id'], who)
    return dict(row)

@router.post('/answers/{aid}/audio')
async def audio(aid: str, request: Request, audio: UploadFile = File(...), chunk_id: str = Form(...)):
    import re
    from app import _process_chunk
    with store.get() as db:
        answer = answer_access(db, aid, owner(request))
        if answer['status'] != 'recording': raise HTTPException(409, '回答已经结束，不能继续上传')
    if not re.fullmatch(r'[A-Za-z0-9_-]{8,64}', chunk_id): raise HTTPException(422, '分片编号无效')
    raw = await audio.read(config.MAX_AUDIO_BYTES + 1)
    if not raw or len(raw) > config.MAX_AUDIO_BYTES: raise HTTPException(413, '音频为空或超过上传限制')
    seq = await run_in_threadpool(_process_chunk, answer['session_id'], raw, audio.filename or 'chunk.webm', chunk_id)
    return {'seq': seq, 'saved': True}

@router.get('/answers/{aid}/recording')
def recording(aid: str, request: Request):
    from app import _recording_file
    with store.get() as db: answer = answer_access(db, aid, owner(request))
    return FileResponse(_recording_file(answer['session_id']), media_type='audio/wav', headers={'Cache-Control': 'private, no-store'})

@router.get('/projects/{pid}/events')
async def events(pid: str, request: Request):
    who = owner(request)
    with store.get() as db: d.owned(db, pid, who)
    async def stream():
        previous = None
        ticks = 0
        while not await request.is_disconnected():
            # Revalidate the employee session on long-lived connections too.
            import employee_auth
            try:
                await employee_auth.current_user(request)
                with store.get() as db: state = d.snapshot(db, pid, who)
            except HTTPException:
                yield 'event: unavailable\ndata: {}\n\n'; return
            encoded = d.dump(state)
            if encoded != previous:
                yield 'event: snapshot\ndata: ' + encoded + '\n\n'; previous = encoded
            ticks += 1
            if ticks % 20 == 0: yield ': keepalive\n\n'
            await asyncio.sleep(1)
    return StreamingResponse(stream(), media_type='text/event-stream', headers={'Cache-Control': 'private, no-store', 'X-Accel-Buffering': 'no'})


def extract_material(raw, filename):
    suffix = Path(filename).suffix.lower()
    if suffix not in {'.txt','.md','.docx','.pdf','.pptx'}: raise HTTPException(422, '支持 TXT、MD、DOCX、PDF 和 PPTX')
    try:
        if suffix in {'.docx','.pptx'}:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                if sum(i.file_size for i in archive.infolist()) > 50 * 1024 * 1024: raise ValueError('expanded_archive_limit')
        if suffix == '.pdf':
            from pypdf import PdfReader
            reader = PdfReader(io.BytesIO(raw))
            if reader.is_encrypted: raise ValueError('encrypted_pdf')
            if len(reader.pages) > 300: raise ValueError('page_limit')
            content = '\n\n'.join(f'[第 {i+1} 页]\n{page.extract_text() or ""}' for i,page in enumerate(reader.pages))
            if not any((page.extract_text() or '').strip() for page in reader.pages): raise HTTPException(422, 'PDF 未提取到正文，可能是扫描件；请先 OCR 或粘贴正文')
        elif suffix == '.pptx':
            from pptx import Presentation
            presentation = Presentation(io.BytesIO(raw))
            content = '\n\n'.join(f'[第 {i+1} 张幻灯片]\n' + '\n'.join(s.text for s in slide.shapes if s.has_text_frame) for i,slide in enumerate(presentation.slides))
        else:
            with tempfile.TemporaryDirectory(prefix='interview-material-') as folder:
                path = Path(folder) / ('material' + suffix); path.write_bytes(raw)
                content = '\n'.join(r['text'] for r in importer.extract_utterances(path))
    except HTTPException: raise
    except Exception as exc: raise HTTPException(422, '材料无法解析，请检查格式、加密状态或改为粘贴正文') from exc
    if not content.strip() or len(content) > 50000: raise HTTPException(422, '未提取到有效正文或超过单份 50000 字限制，请拆分材料')
    return content

@router.post('/projects/{pid}/materials')
async def material(pid: str, request: Request, file: UploadFile = File(...), revision: int = Form(...), request_id: str = Form(...), kind: str = Form('background')):
    with store.get() as db: d.owned(db, pid, owner(request))
    raw = await file.read(20 * 1024 * 1024 + 1)
    if len(raw) > 20 * 1024 * 1024: raise HTTPException(413, '材料文件最多 20MB')
    content = await run_in_threadpool(extract_material, raw, file.filename or '')
    body = Command(action='material', revision=revision, request_id=request_id, title=(Path(file.filename or '材料').name)[:120], content=content, kind=kind)
    return await run_in_threadpool(d.command, pid, owner(request), body.model_dump())
